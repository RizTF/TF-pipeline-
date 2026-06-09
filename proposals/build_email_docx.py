#!/usr/bin/env python3
"""Generate the revised Credico email as a Word document."""
from docx import Document
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x0D, 0x1B, 0x2A)
TEAL = RGBColor(0x00, 0x78, 0x6A)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = TEAL


def para(text="", bold=False):
    p = doc.add_paragraph()
    if text:
        p.add_run(text).bold = bold
    return p


def bullet_rich(lead, rest):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(lead).bold = True
    if rest:
        p.add_run(rest)


# Subject
p = doc.add_paragraph()
p.add_run("Subject: ").bold = True
p.add_run("Your revised advertising proposal — Talent Finder")
doc.add_paragraph()

para("Hi Alicia,")
para("Thank you for your time on our call.")
para("By way of an update: we approached Total Jobs about managing your account on your "
     "behalf, but as you're currently in a direct contract with them, they aren't able to "
     "discuss it with us. Not a problem — our revised package still gives you Total Jobs "
     "exposure as part of a much wider network.")
para("I've reworked the proposal around everything we discussed:")

heading("Job Board Reach")
para("Every advert is published across our full network — CV Library, Total Jobs, Indeed (PPC "
     "sponsored), Jobsite, Talent.com, Job Flurry, Adzuna, Facebook Jobs, Google for Jobs, "
     "CareerJet, Find A Job, ZipRecruiter, Trovit, Zoek, Glassdoor, Job Rapido and X — plus "
     "industry-specific sites, micro-boards and aggregators. In total, exposure across more "
     "than 200 job boards from a single advert.")

heading("Included as Standard")
bullet_rich("Free Applicant Tracking System (ATS)", " — manage every application in one place, "
            "with unlimited user logins and CV forwarding")
bullet_rich("Talent Pool", " — ongoing access to expired campaigns and previous candidate "
            "applications")
bullet_rich("Campaign monitoring", " — we actively manage your adverts to drive the strongest "
            "response, and can update them while they're live")
bullet_rich("Content writing team", " — full support preparing and optimising every advert")
bullet_rich("Dedicated account manager", " — customer service, recruitment strategy, market "
            "analysis and support across all your roles")
bullet_rich("Free email marketing campaigns", "")
bullet_rich("'Invite to Apply'", " — proactively inviting qualified candidates to apply for "
            "your roles")

heading("Pricing Options")
p = doc.add_paragraph()
p.add_run("Our advert credits never expire — buying in a package reduces your cost per "
          "advert.").italic = True
bullet_rich("50 × Premium Advert", " — £255 + VAT per advert")
bullet_rich("100 × Premium Branded Advert", " — £225 + VAT per advert")
para("You're currently paying £275 + VAT per vacancy, so you'll save £20 + VAT per vacancy on "
     "the 50-pack and £50 + VAT per vacancy on the 100-pack — with broader reach than your "
     "current single-board setup.")
para("We'd genuinely like to work with you on a longer-term basis, and we're happy to "
     "negotiate terms to make sure the service is the right fit for Credico.")

heading("Optional Add-On Services")
bullet_rich("AI CV Screening", " — reviews your applications and flags each as suitable, maybe "
            "or unsuitable — £45 + VAT")
bullet_rich("AI CV Screening + Automated Interview Booking (combined)", " — screens "
            "applications and automatically books interviews for qualified candidates scoring "
            "65%+ — £75 + VAT")
bullet_rich("CV Database Search", " — we search two of the UK's largest CV platforms and supply "
            "you with matching CVs — £199 + VAT")
bullet_rich("Candidate Management", " — filtering, telephone screening, shortlisting and "
            "arranging interviews on your behalf — £499 + VAT")

para("I hope this is much better aligned with what you're looking for. Any questions at all, "
     "just let me know — I'm happy to jump on a quick call to walk through it.")
doc.add_paragraph()
para("Best regards,")
doc.add_paragraph()
para("Riz", bold=True)
para("Talent Finder")

out = "/home/user/TF-pipeline-/proposals/Credico_Revised_Proposal_Email.docx"
doc.save(out)
print("Saved", out)
