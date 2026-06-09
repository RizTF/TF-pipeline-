#!/usr/bin/env python3
"""Generate the Credico proposal as a Word document from the agreed content."""
from docx import Document
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x0D, 0x1B, 0x2A)
TEAL = RGBColor(0x00, 0x78, 0x6A)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = TEAL


def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def bullet_rich(lead, rest):
    """Bullet with a bold lead phrase then normal text."""
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(lead).bold = True
    if rest:
        p.add_run(rest)


# Title
t = doc.add_paragraph()
tr = t.add_run("MULTI-BOARD JOB ADVERTISING")
tr.bold = True
tr.font.size = Pt(20)
tr.font.color.rgb = NAVY
t2 = doc.add_paragraph()
t2r = t2.add_run("Proposal for Credico")
t2r.bold = True
t2r.font.size = Pt(14)
t2r.font.color.rgb = TEAL

doc.add_paragraph()
for label, val in [("Prepared for:", "Credico Marketing Limited"),
                   ("Attention:", "Alicia Harris"),
                   ("Prepared by:", "Talent Finder")]:
    p = doc.add_paragraph()
    p.add_run(label + " ").bold = True
    p.add_run(val)

doc.add_paragraph("_" * 60)

# Overview
heading("Overview")
para("Dear Alicia,")
para("Thank you for the opportunity to put this together. This proposal sets out a "
     "like-for-like service to the multi job board advertising Credico currently runs through "
     "Talent Finder — but structured to scale cleanly with your hiring, give you complete "
     "visibility, and take as much of the day-to-day workload off your team as you'd like.")
para("The principle is simple: one credit, one role, advertised across our full job board "
     "network — with everything from advert creation to applicant management handled under a "
     "single account and a single, predictable price.")

# Value of reach
heading("The Value of Multi-Board Reach")
para("Posting a role on one board reaches one audience. Every Talent Finder credit places your "
     "role across four major job boards at once — so each vacancy is seen by a far wider pool "
     "of active candidates from day one:")
for b in ["Indeed", "CV Library", "Reed", "Find a Job"]:
    doc.add_paragraph(b, style="List Bullet")
para("For high-volume, multi-region hiring, that breadth is what fills roles faster and keeps "
     "your pipeline full — without you managing multiple board contracts, logins and invoices.")
p = doc.add_paragraph()
p.add_run("Job boards can be added or removed — this is a guide.").italic = True

# Inclusions
heading("What Each Credit Includes — and Why It Matters")
bullet_rich("Job advert creation", " — professionally written and optimised, so each role "
            "stands out and attracts the right applicants")
bullet_rich("Job posting", " — published and managed across the full network on your behalf")
bullet_rich("Recruitment ATS, unlimited hiring managers", " — add as many of your team as you "
            "need; everyone works from one place with full visibility of every applicant, at no "
            "extra per-seat cost")
bullet_rich("Full account management support", " — a dedicated point of contact who knows your "
            "roles and your business")
bullet_rich("Priority admin support", " — fast turnaround whenever you need to move quickly")
bullet_rich("Advice on salary and positioning", " — guidance to keep each role competitive and "
            "convert more of the right applicants")

# Pricing
heading("Pricing — Transparent, and Scaling With You")
para("The more you advertise, the lower your cost per credit — so your buying power grows as "
     "Credico grows. One fixed price per credit keeps budgeting simple and removes the "
     "unpredictability of pay-as-you-go board spend.")
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
for i, h in enumerate(["Credits", "Price per credit", "Total"]):
    hdr[i].paragraphs[0].add_run(h).bold = True
for credits, price, total in [("100 credits", "£170", "£17,000"),
                              ("300 credits", "£136", "£40,800"),
                              ("500 credits", "£122", "£61,000"),
                              ("800 credits", "£102", "£81,600")]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text, cells[2].text = credits, price, total
p = doc.add_paragraph()
p.add_run("All prices exclude VAT.").italic = True

# Add-on
heading("Optional Add-On — Automated Screening & Interview Scheduling")
para("This is where the real time and cost savings come in. For an additional £75 + VAT per "
     "role:")
bullet_rich("Every applicant is automatically screened", " against your required criteria — so "
            "your team only ever sees candidates that genuinely fit")
bullet_rich("Interviews are automatically scheduled", " around each hiring manager's "
            "availability — no back-and-forth, no diary admin")
para("The result: hours of manual sifting and coordination removed from every role, your "
     "hiring managers focused on interviewing rather than admin, and a faster, more "
     "professional experience for every candidate. At Credico's scale, this is a direct "
     "reduction in staffing cost and time-to-hire.")

# Partnership
heading("A Partnership, Tailored to You")
para("This proposal is a starting point. Once we fully understand Credico's actual "
     "requirements — volumes, regions, role types and how your offices operate — we'll tailor "
     "everything around it. Job boards can be added or removed, volumes adjusted, and the "
     "screening service switched on wherever it adds the most value. You only ever pay for what "
     "fits.")

# Why
heading("Why Talent Finder")
para("Talent Finder has been established for 10 years, with the technological infrastructure to "
     "support Credico's growth — from multi-board advertising and a full applicant tracking "
     "system to automated screening and interview scheduling, all under one roof. As your "
     "hiring scales, the platform and the team scale with you.")

# Next steps
heading("Next Steps")
para("I'd welcome a short call to walk through this, answer any questions, and shape the "
     "package around exactly what Credico needs. Whenever suits you best.")

doc.add_paragraph("_" * 60)
doc.add_paragraph()
para("Best regards,")
doc.add_paragraph()
para("Riz", bold=True)
para("Talent Finder")

out = "/home/user/TF-pipeline-/proposals/Credico_Multi-Board_Advertising_Proposal.docx"
doc.save(out)
print("Saved", out)
